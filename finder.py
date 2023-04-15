import os
import re
import argparse
import chardet
from tqdm import tqdm
import PyPDF2
from docx import Document
import openpyxl

def contains_salary_info(file_path):
    salary_pattern = re.compile(r'(\$|€|£)?\s*\d{1,3}(?:,\d{3})*(?:\.\d{2})?\s*(?:salary|salaries|wage|wages)|(?:salary|salaries|wage|wages)\s*\d{1,3}(?:,\d{3})*(?:\.\d{2})?\s*(\$|€|£)?', re.IGNORECASE)
    
    if file_path.endswith('.pdf'):
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfFileReader(file)
                file_content = ' '.join([pdf_reader.getPage(i).extractText() for i in range(pdf_reader.getNumPages())])
        except:
            return False
    elif file_path.endswith('.docx'):
        try:
            doc = Document(file_path)
            file_content = ' '.join([para.text for para in doc.paragraphs])
        except:
            return False
    elif file_path.endswith('.xlsx'):
        try:
            workbook = openpyxl.load_workbook(file_path)
            file_content = ' '.join([' '.join([str(cell.value) for cell in row]) for sheet in workbook for row in sheet.iter_rows()])
        except:
            return False
    else:
        with open(file_path, 'rb') as file:
            file_content = file.read()
        encoding = chardet.detect(file_content)['encoding']
        file_content = file_content.decode(encoding)

    if salary_pattern.search(file_content):
        return True
    return False

def find_salary_files(directory):
    salary_files = []

    for root, _, files in os.walk(directory):
        for file in tqdm(files, desc="Searching for salary files"):
            file_path = os.path.join(root, file)
            try:
                if contains_salary_info(file_path):
                    salary_files.append(file_path)
            except (UnicodeDecodeError, LookupError):
                print(f"Cannot read the file: {file_path}")

    return salary_files

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Search for files containing salary information.")
    parser.add_argument("directory", type=str, help="Directory to search for salary files.")
    args = parser.parse_args()

    if os.path.isdir(args.directory):
        salary_files = find_salary_files(args.directory)
        if salary_files:
            print("Salary files found:")
            for file in salary_files:
                print(f"- {file}")
        else:
            print("No salary files found.")
    else:
        print("Invalid directory.")
